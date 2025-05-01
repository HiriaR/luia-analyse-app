
import streamlit as st
import pandas as pd
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches
from datetime import datetime
import tempfile
import os

LUIA_PRICE_USD = 0.00010411

st.set_page_config(page_title="Analyseur Wallets LUIA by Luia Team", layout="wide")
st.title("ðŸ“Š Analyse complÃ¨te des Wallets du Token LUIA")

tx_file = st.file_uploader("ðŸ§¾ Fichier de transactions (CSV)", type="csv")
holders_file = st.file_uploader("ðŸ“„ Fichier de holders (CSV)", type="csv")

st.subheader("ðŸ”– Associer des noms personnalisÃ©s Ã  des wallets")
raw_aliases = st.text_area("Noms (ex: 0xabc123: Kale le crÃ©ateur)", height=150)

alias_map = {}
if raw_aliases:
    for line in raw_aliases.splitlines():
        if ":" in line:
            addr, name = line.split(":", 1)
            alias_map[addr.strip().lower()] = name.strip()

if tx_file and holders_file:
    st.success("âœ… Fichiers chargÃ©s, traitement en cours...")

    transactions = pd.read_csv(tx_file)
    holders = pd.read_csv(holders_file)

    transactions['DateTime (UTC)'] = pd.to_datetime(transactions['DateTime (UTC)'])
    transactions['Quantity'] = transactions['Quantity'].astype(str).str.replace(',', '').astype(float)
    holders['Balance'] = holders['Balance'].astype(str).str.replace(',', '').astype(float)

    transactions['From'] = transactions['From'].astype(str).str.strip().str.lower()
    holders['HolderAddress'] = holders['HolderAddress'].astype(str).str.strip().str.lower()

    latest_time = transactions['DateTime (UTC)'].max()
    start_time = latest_time - pd.Timedelta(hours=24)
    recent_sales = transactions[transactions['DateTime (UTC)'] >= start_time]
    sold_by_wallet = recent_sales.groupby('From')['Quantity'].sum().reset_index()
    sold_by_wallet.columns = ['Wallet', 'Total_Vendu_24h']

    top_20 = holders.head(20)
    merged = top_20.merge(sold_by_wallet, left_on='HolderAddress', right_on='Wallet', how='left')
    exclude_wallets = ["0x2200c5ac68f2b7ed93f2dfda39d8fdd2eddfddf6"]
    merged = merged[~merged['HolderAddress'].isin([e.lower() for e in exclude_wallets])]

    merged['Nom'] = merged['HolderAddress'].map(alias_map).fillna("")
    merged['Label'] = merged['Nom']
    merged.loc[merged['Label'] == '', 'Label'] = merged['HolderAddress']

    merged['Solde_USD'] = merged['Balance'] * LUIA_PRICE_USD
    merged['Ventes_USD'] = merged['Total_Vendu_24h'].fillna(0) * LUIA_PRICE_USD

    st.subheader("ðŸ“‹ RÃ©sumÃ© des top 20 holders")
    st.dataframe(merged[['Nom', 'HolderAddress', 'Balance', 'Solde_USD', 'Total_Vendu_24h', 'Ventes_USD']])

    st.subheader("ðŸ“ˆ Graphique des soldes vs ventes")
    fig, ax = plt.subplots(figsize=(10, 5))
    ax.bar(merged['Label'], merged['Balance'], label='Solde actuel')
    ax.bar(merged['Label'], merged['Total_Vendu_24h'].fillna(0), label='Vendu (24h)', alpha=0.7)
    ax.set_xticks(range(len(merged)))
    ax.set_xticklabels(merged['Label'], rotation=90)
    ax.legend()
    st.pyplot(fig)

    st.subheader("ðŸ“¤ GÃ©nÃ©ration du rapport Word")

    if st.button("ðŸ“ TÃ©lÃ©charger le rapport complet Word"):
        doc = Document()
        doc.add_heading("Analyse des ventes sur 24h â€“ Token LUIA", level=1)
        doc.add_paragraph(f"Date : {datetime.utcnow().strftime('%d/%m/%Y %H:%M UTC')}")
        doc.add_paragraph("Analyse des mouvements des top holders, avec Ã©quivalents en USD, dÃ©tails des ventes et graphique.")

        tmp_chart = tempfile.mktemp(suffix=".png")
        fig.savefig(tmp_chart)
        doc.add_picture(tmp_chart, width=Inches(6.5))

        doc.add_heading("Top 20 Holders (hors PancakeSwap)", level=2)
        table = doc.add_table(rows=1, cols=6)
        table.style = 'Light Grid'
        hdr = table.rows[0].cells
        hdr[0].text = "Nom"
        hdr[1].text = "Wallet"
        hdr[2].text = "Solde (LUIA)"
        hdr[3].text = "Solde (USD)"
        hdr[4].text = "Ventes 24h (LUIA)"
        hdr[5].text = "Ventes 24h (USD)"

        for _, row in merged.iterrows():
            cells = table.add_row().cells
            cells[0].text = row['Nom']
            cells[1].text = row['HolderAddress']
            cells[2].text = f"{row['Balance']:,.2f}"
            cells[3].text = f"{row['Solde_USD']:,.2f} $"
            cells[4].text = f"{row['Total_Vendu_24h'] if pd.notna(row['Total_Vendu_24h']) else 0:,.2f}"
            cells[5].text = f"{row['Ventes_USD']:,.2f} $"

        doc.add_page_break()
        doc.add_heading("Top 5 vendeurs â€“ DÃ©tail des transactions", level=2)
        top5 = merged.nlargest(5, 'Total_Vendu_24h')
        for _, wallet in top5.iterrows():
            addr = wallet['HolderAddress']
            name = wallet['Nom'] or addr
            doc.add_heading(f"{name}", level=3)
            subtx = recent_sales[recent_sales['From'].astype(str).str.strip().str.lower() == addr]
            subtx = subtx.sort_values('DateTime (UTC)')
            if subtx.empty:
                doc.add_paragraph("Aucune transaction trouvÃ©e.")
            else:
                tx_table = doc.add_table(rows=1, cols=3)
                tx_table.style = 'Table Grid'
                hdr = tx_table.rows[0].cells
                hdr[0].text = "Date"
                hdr[1].text = "Montant (LUIA)"
                hdr[2].text = "Montant (USD)"
                for _, tx in subtx.iterrows():
                    row = tx_table.add_row().cells
                    row[0].text = tx['DateTime (UTC)'].strftime('%d/%m %H:%M')
                    row[1].text = f"{tx['Quantity']:,.2f}"
                    row[2].text = f"{tx['Quantity'] * LUIA_PRICE_USD:,.2f} $"

        doc.add_paragraph("\n\nRapport gÃ©nÃ©rÃ© automatiquement et signÃ© par Jarvis.")
        output = tempfile.mktemp(suffix=".docx")
        doc.save(output)

        with open(output, "rb") as f:
            st.download_button("ðŸ“„ TÃ©lÃ©charger le rapport Word", f, file_name="rapport_luia_complet.docx")

        os.remove(tmp_chart)
        os.remove(output)